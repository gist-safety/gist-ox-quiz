# 엑셀 → questions.js 변환
#  - 문제: ..\퀴즈 엑셀 파일\OX문제 선별.xlsx 의 'Sheet1' (번호 칸에 쉬움/어려움)
#  - 영문: ..\퀴즈 엑셀 파일\OX퀴즈_영문번역.xlsx (한글 문제 문장으로 연결)
# 사용법: 문제변환.bat 실행
import json, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_DIR = os.path.join(HERE, '..', '퀴즈 엑셀 파일')
KO_XLSX = os.path.join(XLSX_DIR, 'OX문제 선별.xlsx')
EN_XLSX = os.path.join(XLSX_DIR, 'OX퀴즈_영문번역.xlsx')

def s(v):
    return str(v or '').strip()

en = {}
if os.path.exists(EN_XLSX):
    for r in openpyxl.load_workbook(EN_XLSX, data_only=True).active.iter_rows(min_row=2, values_only=True):
        if r and r[1]:
            en[s(r[1])] = {'category': s(r[3]), 'question': s(r[4]), 'explanation': s(r[5])}

ws = openpyxl.load_workbook(KO_XLSX, data_only=True)['Sheet1']
quiz = {'easy': [], 'hard': []}
missing = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r or not r[2] or s(r[3]).upper() not in ('O', 'X'):
        continue
    level = 'easy' if s(r[0]) == '쉬움' else 'hard' if s(r[0]) == '어려움' else None
    if not level:
        continue
    q = s(r[2])
    item = {
        'answer': s(r[3]).upper(),
        'ko': {'category': s(r[1]), 'question': q, 'explanation': s(r[4])},
    }
    if q in en and en[q]['question']:
        item['en'] = en[q]
    else:
        missing.append(q)
    quiz[level].append(item)

with open(os.path.join(HERE, 'questions.js'), 'w', encoding='utf-8') as f:
    f.write('// 자동 생성 파일입니다. 문제를 바꾸려면 엑셀을 고친 뒤 문제변환.bat 을 실행하세요.\n')
    f.write('window.QUIZ = ' + json.dumps(quiz, ensure_ascii=False, indent=1) + ';\n')

print(f"완료: 쉬움 {len(quiz['easy'])}문항, 어려움 {len(quiz['hard'])}문항 → questions.js")
if missing:
    print(f'※ 영문 번역이 없는 문항 {len(missing)}개 (영어 모드에서 한글로 표시됨):')
    for q in missing:
        print('   -', q)
print('※ 태블릿 앱에 반영하려면 수정한 파일을 다시 배포해야 합니다.')
